"""Assemble the NestQuest final-year project report into a formatted .docx.

Reads docs/chapters/*.md (constrained markdown), embeds docs/figures/*.png and
static/images/*.png, and emits docs/NestQuest_Final_Year_Project_Report.docx with
academic formatting, a Word TOC / List of Figures / List of Tables (auto-updating
fields), chapter-wise headings, tables, code blocks and annexures.

Run:  python docs/build_report.py
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from PIL import Image
except Exception:  # pragma: no cover
    Image = None

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG_DIR = os.path.join(HERE, "figures")
IMG_DIR = os.path.join(ROOT, "static", "images")
CH_DIR = os.path.join(HERE, "chapters")
OUT = os.path.join(HERE, "NestQuest_Final_Year_Project_Report.docx")

TNR = "Times New Roman"
MONO = "Consolas"
USABLE_WIN = 5.0  # inches of usable page width used for figures (A4 minus margins, tightened)
INK = RGBColor(0x11, 0x11, 0x11)


# ----------------------------------------------------------------- low-level XML helpers
def _set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _set_table_borders(table, color="BBBBBB", sz="6"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _field(paragraph, instr, placeholder=None):
    """Insert a Word field (e.g. TOC/PAGE) into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(begin)
    run._r.append(instr_el)
    run._r.append(sep)
    if placeholder:
        paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _tc_entry(paragraph, text, flag):
    """Hidden TC field so List of Figures/Tables auto-populate with page numbers."""
    safe = text.replace('"', "'")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TC "{safe}" \\f {flag} '
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _page_num_format(section, fmt, start=1):
    sectPr = section._sectPr
    for el in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(el)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def _footer_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(p, " PAGE ")


# ----------------------------------------------------------------- styles / page setup
def setup_document():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), TNR)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)

    sizes = {1: 16, 2: 14, 3: 12, 4: 12}
    for lvl, sz in sizes.items():
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = TNR
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.italic = (lvl == 4)
        st.font.color.rgb = INK
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before = Pt(8 if lvl <= 2 else 6)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.left_margin = Inches(1.5)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    return doc


def _update_fields_on_open(doc):
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


# ----------------------------------------------------------------- content primitives
def center(doc, text, size=12, bold=False, italic=False, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def add_inline(paragraph, text):
    for tok in re.split(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = MONO
            r.font.size = Pt(10.5)
        else:
            paragraph.add_run(tok)


def add_code_block(doc, lines):
    while lines and lines[0].strip() == "":
        lines.pop(0)
    while lines and lines[-1].strip() == "":
        lines.pop()
    if not lines:
        return
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _set_cell_shading(cell, "F3F4F6")
    _set_table_borders(tbl, color="D8DCE4", sz="4")
    first = True
    for ln in lines:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        r = para.add_run(ln.replace("\t", "    "))
        r.font.name = MONO
        r.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows_md, caption):
    rows = []
    for ln in rows_md:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    # drop the |---| separator row
    rows = [r for r in rows if not all(set(c) <= set("-: ") and c for c in r)]
    if not rows:
        return
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(2)
        r = cap.add_run(caption)
        r.bold = True
        r.font.size = Pt(10)
        _tc_entry(cap, caption, "T")
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            txt = row[ci] if ci < len(row) else ""
            add_inline(cell.paragraphs[0], txt)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9.5)
                if ri == 0:
                    run.bold = True
            if ri == 0:
                _set_cell_shading(cell, "1F4B99")
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc, path, caption):
    if not os.path.exists(path):
        center(doc, f"[Missing figure: {os.path.basename(path)}]", size=10, italic=True)
        return
    width = USABLE_WIN
    if Image is not None:
        try:
            with Image.open(path) as im:
                w_px, h_px = im.size
            win = w_px / 150.0
            width = min(USABLE_WIN, win)
            # keep tall figures from overflowing the page height
            if (h_px / w_px) * width > 6.4:
                width = 6.4 * (w_px / h_px)
                width = min(width, USABLE_WIN)
        except Exception:
            pass
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(10)
    _tc_entry(cap, caption, "F")


def add_screenshot_placeholder(doc, caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _set_cell_shading(cell, "EEF1F6")
    _set_table_borders(tbl, color="C2C9D6", sz="6")
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(10)
    r = para.add_run("[ Screenshot placeholder — capture from the running application ]")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x60, 0x7A)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    rr = cap.add_run(caption)
    rr.italic = True
    rr.font.size = Pt(10)
    _tc_entry(cap, caption, "F")


def chapter_heading(doc, text):
    doc.add_heading(text, level=1)


# ----------------------------------------------------------------- markdown renderer
def render_markdown(doc, text):
    lines = text.replace("\r\n", "\n").split("\n")
    i, n = 0, len(lines)
    pending_caption = None
    buf = []

    def flush():
        nonlocal buf
        if buf:
            t = " ".join(x.strip() for x in buf).strip()
            if t:
                add_inline(doc.add_paragraph(), t)
            buf = []

    while i < n:
        s = lines[i].strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code)
            continue
        if s.startswith("|") and s.count("|") >= 2:
            flush()
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            add_table(doc, tbl, pending_caption)
            pending_caption = None
            continue
        if s.startswith("[[TABLE"):
            flush()
            inner = s[2:-2].strip()[len("TABLE"):].strip()
            pending_caption = inner
            i += 1
            continue
        if s.startswith("[[FIG") or s.startswith("[[EDA"):
            flush()
            inner = s[2:-2].strip()
            kind, rest = inner.split(" ", 1)
            fname, caption = [x.strip() for x in rest.split("|", 1)]
            folder = FIG_DIR if kind == "FIG" else IMG_DIR
            add_figure(doc, os.path.join(folder, fname), caption)
            i += 1
            continue
        if s.startswith("[[SHOT"):
            flush()
            inner = s[2:-2].strip()
            caption = inner.split("|", 1)[1].strip()
            add_screenshot_placeholder(doc, caption)
            i += 1
            continue
        if s == "[[PAGEBREAK]]":
            flush()
            doc.add_page_break()
            i += 1
            continue
        if s.startswith("#### "):
            flush(); doc.add_heading(s[5:].strip(), level=4); i += 1; continue
        if s.startswith("### "):
            flush(); doc.add_heading(s[4:].strip(), level=3); i += 1; continue
        if s.startswith("## "):
            flush(); doc.add_heading(s[3:].strip(), level=2); i += 1; continue
        if s.startswith("# "):
            flush(); chapter_heading(doc, s[2:].strip()); i += 1; continue
        if s.startswith("- "):
            flush(); add_inline(doc.add_paragraph(style="List Bullet"), s[2:].strip()); i += 1; continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:
            flush(); add_inline(doc.add_paragraph(style="List Number"), m.group(2).strip()); i += 1; continue
        if s == "" or s == "---":
            flush(); i += 1; continue
        buf.append(s)
        i += 1
    flush()


def read(path):
    with open(path, "r", encoding="utf-8") as fh:
        return fh.read()


# ----------------------------------------------------------------- front matter
PH = "[TO BE FILLED]"


def title_page(doc):
    center(doc, "NESTQUEST", size=26, bold=True, space_before=24, space_after=2)
    center(doc, "A Brokerage-Free Property Marketplace with an AI Price Estimator",
           size=14, bold=True, space_after=18)
    center(doc, "A PROJECT REPORT", size=13, bold=True, space_after=2)
    center(doc, "Submitted in partial fulfilment of the requirements for the award of the degree of",
           size=12, space_after=2)
    center(doc, "[Degree / Course — e.g. Bachelor of Computer Applications]", size=12, bold=True, space_after=18)
    center(doc, "Submitted by", size=12, space_after=2)
    center(doc, "[Student Name]   ([Register Number])", size=12, bold=True, space_after=18)
    center(doc, "Under the guidance of", size=12, space_after=2)
    center(doc, "[Guide Name], [Designation]", size=12, bold=True, space_after=18)
    center(doc, "[Department Name]", size=12, bold=True, space_after=2)
    center(doc, "[College / Institution Name]", size=13, bold=True, space_after=2)
    center(doc, "[University Name]", size=12, space_after=2)
    center(doc, "[Academic Year]", size=12, bold=True, space_after=2)
    doc.add_page_break()


def certificate(doc):
    doc.add_heading("BONAFIDE CERTIFICATE", level=1)
    for t in [
        "This is to certify that the project report entitled “NestQuest — A Brokerage-Free "
        "Property Marketplace with an AI Price Estimator” is a bonafide record of the work carried "
        "out by [Student Name] (Register Number [Register Number]) in partial fulfilment of the "
        "requirements for the award of the degree of [Degree / Course] of [University Name], during the "
        "academic year [Academic Year].",
        "The work presented in this report is original and has not been submitted earlier, in part or in "
        "full, for the award of any other degree or diploma of this or any other institution.",
    ]:
        add_inline(doc.add_paragraph(), t)
    doc.add_paragraph()
    doc.add_paragraph()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).paragraphs[0].add_run("Signature of the Guide\n[Guide Name]\n[Designation]")
    tbl.cell(0, 1).paragraphs[0].add_run("Signature of the Head of Department\n[HOD Name]")
    doc.add_paragraph()
    add_inline(doc.add_paragraph(),
               "Submitted for the project viva-voce examination held on ____________________.")
    doc.add_paragraph()
    tbl2 = doc.add_table(rows=1, cols=2)
    tbl2.cell(0, 0).paragraphs[0].add_run("Internal Examiner")
    tbl2.cell(0, 1).paragraphs[0].add_run("External Examiner")
    doc.add_page_break()


def declaration(doc):
    doc.add_heading("DECLARATION", level=1)
    for t in [
        "I, [Student Name], bearing Register Number [Register Number], hereby declare that the project "
        "report entitled “NestQuest — A Brokerage-Free Property Marketplace with an AI Price "
        "Estimator” submitted to [University Name] is a record of original work done by me under the "
        "guidance of [Guide Name], [Designation], [Department Name], [College / Institution Name].",
        "I further declare that this report has not been submitted, in whole or in part, for the award of "
        "any degree, diploma or other similar title to any other university or institution.",
    ]:
        add_inline(doc.add_paragraph(), t)
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Place: [City]\nDate:  ____________\n\n\t\t\t\t\t\t[Student Name]")
    doc.add_page_break()


def acknowledgement(doc):
    doc.add_heading("ACKNOWLEDGEMENT", level=1)
    for t in [
        "I take this opportunity to express my sincere gratitude to all those who supported me throughout "
        "this project. I am deeply thankful to my guide, [Guide Name], [Designation], for the valuable "
        "guidance, continuous encouragement and constructive feedback that shaped this work.",
        "I extend my thanks to [HOD Name], Head of the [Department Name], and to all the faculty members "
        "of [College / Institution Name] for providing the resources and the academic environment that "
        "made this project possible.",
        "Finally, I am grateful to my family and friends for their constant support and motivation during "
        "the course of this project.",
    ]:
        add_inline(doc.add_paragraph(), t)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("[Student Name]").bold = True
    doc.add_page_break()


def abstract(doc):
    doc.add_heading("ABSTRACT", level=1)
    render_markdown(doc, read(os.path.join(CH_DIR, "abstract.md")))
    doc.add_page_break()


def toc_page(doc, title, instr):
    doc.add_heading(title, level=1)
    p = doc.add_paragraph()
    _field(p, instr, placeholder="Right-click and choose “Update Field” to populate this list "
                                 "(or press Ctrl+A then F9).")
    doc.add_page_break()


ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"),
    ("ASGI", "Asynchronous Server Gateway Interface"), ("BHK", "Bedroom-Hall-Kitchen"),
    ("CART", "Classification and Regression Tree"), ("CDN", "Content Delivery Network"),
    ("CRUD", "Create, Read, Update, Delete"), ("CSRF", "Cross-Site Request Forgery"),
    ("CSS", "Cascading Style Sheets"), ("DFD", "Data-Flow Diagram"),
    ("EDA", "Exploratory Data Analysis"), ("ER", "Entity-Relationship"),
    ("HTML", "HyperText Markup Language"), ("HTTP", "HyperText Transfer Protocol"),
    ("HTTPS", "HTTP Secure"), ("JSON", "JavaScript Object Notation"),
    ("JWT", "JSON Web Token"), ("MAE", "Mean Absolute Error"),
    ("ML", "Machine Learning"), ("ODM", "Object-Document Mapper"),
    ("ReDoS", "Regular-Expression Denial of Service"), ("REST", "Representational State Transfer"),
    ("RMSE", "Root Mean Squared Error"), ("SMTP", "Simple Mail Transfer Protocol"),
    ("SPA", "Single-Page Application"), ("TLS", "Transport Layer Security"),
    ("UAT", "User Acceptance Testing"), ("UI", "User Interface"),
    ("URL", "Uniform Resource Locator"), ("WDAC", "Windows Defender Application Control"),
]


def abbreviations_page(doc):
    doc.add_heading("LIST OF ABBREVIATIONS", level=1)
    table = doc.add_table(rows=1, cols=2)
    _set_table_borders(table)
    hdr = table.rows[0].cells
    for c, t in zip(hdr, ("Abbreviation", "Expansion")):
        c.paragraphs[0].add_run(t).bold = True
        _set_cell_shading(c, "1F4B99")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ab, full in ABBREVIATIONS:
        row = table.add_row().cells
        row[0].paragraphs[0].add_run(ab).bold = True
        row[1].paragraphs[0].add_run(full)
    doc.add_page_break()


# ----------------------------------------------------------------- annexures
def annexure_a(doc):
    chapter_heading(doc, "ANNEXURE A — CODING")
    add_inline(doc.add_paragraph(),
               "This annexure lists representative source modules of the NestQuest implementation. "
               "The complete source tree is available in the project repository; the excerpts below "
               "cover the security layer, the AI price estimator, the Cloudinary upload service and the "
               "corresponding API routers.")
    files = [
        ("A.1  backend/security.py", "backend/security.py", 46),
        ("A.2  backend/ml/predict.py", "backend/ml/predict.py", 48),
        ("A.3  backend/cloudinary_service.py", "backend/cloudinary_service.py", 48),
        ("A.4  backend/routers/uploads.py", "backend/routers/uploads.py", 46),
        ("A.5  backend/routers/ml.py", "backend/routers/ml.py", 42),
    ]
    for heading, rel, maxlines in files:
        doc.add_heading(heading, level=2)
        path = os.path.join(ROOT, rel)
        if os.path.exists(path):
            code = read(path).split("\n")
            truncated = len(code) > maxlines
            if truncated:
                code = code[:maxlines]
            add_code_block(doc, code)
            if truncated:
                note = doc.add_paragraph()
                r = note.add_run("… (excerpt; the complete module is available in the project repository).")
                r.italic = True
                r.font.size = Pt(9)
        else:
            add_inline(doc.add_paragraph(), f"[Source file not found: {rel}]")
    doc.add_page_break()


SCREENSHOTS = [
    "Figure B.1 — Home / landing page with featured listings and hero search.",
    "Figure B.2 — Browse / search results page with filters (city, type, budget, bedrooms).",
    "Figure B.3 — Property detail page with image gallery and enquiry action.",
    "Figure B.4 — Post-property form with the Cloudinary image uploader and AI price check.",
    "Figure B.5 — AI price-estimator (/insights) page with model metrics and EDA charts.",
    "Figure B.6 — User dashboard (listings, shortlist, enquiries received/sent).",
    "Figure B.7 — Admin panel: statistics, listing moderation and user management.",
    "Figure B.8 — Login / signup pages.",
]


def annexure_b(doc):
    chapter_heading(doc, "ANNEXURE B — SCREENSHOTS")
    add_inline(doc.add_paragraph(),
               "The following screens should be captured from the running application "
               "(http://localhost:8000) and inserted here. Placeholders are provided with the exact "
               "screen each figure should show.")
    for cap in SCREENSHOTS:
        add_screenshot_placeholder(doc, cap)
    doc.add_page_break()


REFERENCES = [
    'S. Ramírez, “FastAPI — Modern, Fast Web Framework for Building APIs with Python,” official documentation, https://fastapi.tiangolo.com (accessed 2026).',
    'F. Pedregosa et al., “Scikit-learn: Machine Learning in Python,” Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.',
    'T. Hastie, R. Tibshirani, and J. Friedman, The Elements of Statistical Learning: Data Mining, Inference, and Prediction, 2nd ed. Springer, 2009.',
    'W. McKinney, “Data Structures for Statistical Computing in Python,” in Proc. 9th Python in Science Conf. (SciPy), pp. 51–56, 2010.',
    'C. R. Harris et al., “Array Programming with NumPy,” Nature, vol. 585, pp. 357–362, 2020.',
    'N. Provos and D. Mazières, “A Future-Adaptable Password Scheme,” in Proc. USENIX Annual Technical Conf., 1999.',
    'M. Jones, J. Bradley, and N. Sakimura, “JSON Web Token (JWT),” RFC 7519, IETF, 2015.',
    'OWASP Foundation, “OWASP Top 10 — 2021,” https://owasp.org/Top10/ (accessed 2026).',
    'MongoDB Inc., “MongoDB Manual,” official documentation, https://www.mongodb.com/docs/ (accessed 2026).',
    'A. Ronacher, “Jinja2 Documentation,” Pallets Projects, https://jinja.palletsprojects.com (accessed 2026).',
    'Cloudinary Ltd., “Cloudinary Documentation — Image and Video Upload API,” https://cloudinary.com/documentation (accessed 2026).',
    'NoBroker Technologies, “Brokerage-free property platform,” https://www.nobroker.in (accessed 2026).',
    'Info Edge (India) Ltd., “99acres real-estate portal,” https://www.99acres.com (accessed 2026).',
    'Magicbricks Realty Services, “Magicbricks property portal,” https://www.magicbricks.com (accessed 2026).',
    'Zillow Group, “Zestimate home-valuation methodology,” https://www.zillow.com/z/zestimate/ (accessed 2026).',
    'Pydantic contributors, “Pydantic v2 Documentation — Data Validation for Python,” https://docs.pydantic.dev (accessed 2026).',
]


def annexure_c(doc):
    chapter_heading(doc, "ANNEXURE C — REFERENCES")
    add_inline(doc.add_paragraph(),
               "References are listed in IEEE style. All sources are genuine and publicly traceable.")
    for idx, ref in enumerate(REFERENCES, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        p.add_run(f"[{idx}]\t")
        add_inline(p, ref)


# ----------------------------------------------------------------- assembly
def main():
    doc = setup_document()

    # ----- front matter (section 0) -----
    title_page(doc)
    certificate(doc)
    declaration(doc)
    acknowledgement(doc)
    abstract(doc)
    toc_page(doc, "TABLE OF CONTENTS", ' TOC \\o "1-4" \\h \\z \\u ')
    toc_page(doc, "LIST OF FIGURES", ' TOC \\h \\z \\f F ')
    toc_page(doc, "LIST OF TABLES", ' TOC \\h \\z \\f T ')
    abbreviations_page(doc)

    fm = doc.sections[0]
    fm.different_first_page_header_footer = True
    _page_num_format(fm, "lowerRoman", start=1)
    _footer_page_number(fm)

    # ----- body (section 1) -----
    doc.add_section(WD_SECTION.NEW_PAGE)
    body = doc.sections[1]
    body.different_first_page_header_footer = False
    body.left_margin = Inches(1.5)
    body.right_margin = Inches(1.0)
    body.top_margin = Inches(1.0)
    body.bottom_margin = Inches(1.0)
    _page_num_format(body, "decimal", start=1)
    _footer_page_number(body)

    for idx in range(1, 9):
        if idx > 1:
            doc.add_page_break()
        render_markdown(doc, read(os.path.join(CH_DIR, f"ch{idx}.md")))

    doc.add_page_break()
    annexure_a(doc)
    annexure_b(doc)
    annexure_c(doc)

    _update_fields_on_open(doc)
    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
